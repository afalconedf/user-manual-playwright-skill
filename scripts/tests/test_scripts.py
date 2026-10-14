import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

CHECK = Path(__file__).resolve().parents[1] / "check_manual.py"
MD2DOCX = Path(__file__).resolve().parents[1] / "md_to_docx.py"


def run_check(manual_text, setup=None):
    with tempfile.TemporaryDirectory() as tmp:
        d = Path(tmp)
        if setup:
            setup(d)
        md = d / "manual.md"
        md.write_text(manual_text, encoding="utf-8")
        r = subprocess.run([sys.executable, str(CHECK), str(md)], capture_output=True, text=True, encoding="utf-8")
        return r.returncode, r.stdout


class CheckManualTests(unittest.TestCase):
    def test_prose_password_is_not_a_credential(self):
        code, out = run_check("**Password:** type your password\n")
        self.assertEqual(code, 0, out)
        self.assertNotIn("possible credential", out)

    def test_password_assignment_is_flagged(self):
        code, out = run_check("password=Secr3t!\n")
        self.assertEqual(code, 1, out)
        self.assertIn("possible credential", out)

    def test_token_value_is_flagged(self):
        code, out = run_check("Token: abcd1234\n")
        self.assertEqual(code, 1, out)
        self.assertIn("possible credential", out)

    def test_email_inside_image_target_is_ignored(self):
        def setup(d):
            (d / "images").mkdir()
            Image.new("RGB", (1920, 1080)).save(d / "images" / "a@2x.png")

        code, out = run_check("![x](images/a@2x.png)\n", setup)
        self.assertEqual(code, 0, out)
        self.assertNotIn("e-mail", out)

    def test_email_in_text_is_flagged(self):
        code, out = run_check("contact mario.rossi@corp.it\n")
        self.assertEqual(code, 1, out)
        self.assertIn("e-mail in text", out)

    def test_credentials_flagged(self):
        for line in ["**Password:** Secr3t123", "DB_PASSWORD=Secr3t123", "client_secret=abcd1234", "password: 123"]:
            with self.subTest(line=line):
                code, out = run_check(line + "\n")
                self.assertEqual(code, 1, out)
                self.assertIn("possible credential", out)

    def test_prose_not_flagged(self):
        for line in ["**Password:** type your password", "Password: <your password>", "| Password | the password field |"]:
            with self.subTest(line=line):
                code, out = run_check(line + "\n")
                self.assertEqual(code, 0, out)
                self.assertNotIn("possible credential", out)

    def test_mailto_link_email_flagged(self):
        code, out = run_check("[mario.rossi@corp.it](mailto:mario.rossi@corp.it)\n")
        self.assertEqual(code, 1, out)
        self.assertIn("e-mail in text", out)


def run_md2docx(md_text, setup=None):
    with tempfile.TemporaryDirectory() as tmp:
        d = Path(tmp)
        if setup:
            setup(d)
        md = d / "manual.md"
        md.write_text(md_text, encoding="utf-8")
        out = d / "manual.docx"
        r = subprocess.run([sys.executable, str(MD2DOCX), str(md), "-o", str(out)], capture_output=True, text=True, encoding="utf-8")
        doc = Document(out) if out.exists() else None
        return r.returncode, r.stderr, doc


class MdToDocxTests(unittest.TestCase):
    def test_numbered_list_restarts_after_heading(self):
        code, err, doc = run_md2docx("# Manual\n\n1. first\n\n2. second\n\n## Next\n\n1. again\n")
        self.assertEqual(code, 0, err)
        nums = [p._p.pPr.numPr.numId.val for p in doc.paragraphs if p.style.name == "List Number"]
        self.assertEqual(len(nums), 3)
        self.assertEqual(nums[0], nums[1])
        self.assertNotEqual(nums[2], nums[0])
        num = doc.part.numbering_part.element.num_having_numId(nums[2])
        self.assertEqual(num.find(".//" + qn("w:startOverride")).get(qn("w:val")), "1")

    def test_list_restart_takes_start_number_from_item(self):
        def setup(d):
            (d / "images").mkdir()
            Image.new("RGB", (1920, 1080)).save(d / "images" / "s.png")

        code, err, doc = run_md2docx("1. a\n![s](images/s.png)\n2. b\n", setup)
        self.assertEqual(code, 0, err)
        b = next(p for p in doc.paragraphs if p.text == "b")
        num = doc.part.numbering_part.element.num_having_numId(b._p.pPr.numPr.numId.val)
        self.assertEqual(num.find(".//" + qn("w:startOverride")).get(qn("w:val")), "2")

    def test_missing_image_fails_run(self):
        code, err, doc = run_md2docx("![pic](images/missing.png)\n")
        self.assertEqual(code, 1, err)
        self.assertIn("missing image", err)
        self.assertIsNotNone(doc)

    def test_table_image_resolves_from_markdown_folder(self):
        def setup(d):
            (d / "images").mkdir()
            Image.new("RGB", (10, 10)).save(d / "images" / "a.png")

        code, err, doc = run_md2docx("| H |\n|---|\n| ![a](images/a.png) |\n", setup)
        self.assertEqual(code, 0, err)
        self.assertEqual(len(doc.inline_shapes), 1)

    def test_escaped_pipe_in_table_cell(self):
        code, err, doc = run_md2docx("| H1 | H2 |\n|---|---|\n| a \\| b | c |\n")
        self.assertEqual(code, 0, err)
        t = doc.tables[0]
        self.assertEqual(len(t.columns), 2)
        self.assertEqual(t.cell(1, 0).text, "a | b")

    def test_link_keeps_url(self):
        code, err, doc = run_md2docx("[portal](https://example.com)\n\n[https://example.com](https://example.com)\n")
        self.assertEqual(code, 0, err)
        texts = [p.text for p in doc.paragraphs]
        self.assertIn("portal (https://example.com)", texts)
        self.assertIn("https://example.com", texts)


if __name__ == "__main__":
    unittest.main()
